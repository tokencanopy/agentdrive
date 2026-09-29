"""Synthetic workbook fixtures, generated at import rather than committed.

No binary blobs in the tree: a committed .xlsx is opaque to review and rots
against the library that wrote it. These build what each test needs from
openpyxl directly, so a reader can see exactly what is in the file.

All content is synthetic (`.test` / `example.com`), per the repository's
never-commit-real-data rule.
"""

from __future__ import annotations

import datetime as dt
import io

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    """`{"Q3": [[...], [...]]}` -> xlsx bytes, sheets in dict order."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def make_xlsx():
    return build_xlsx


@pytest.fixture
def values_workbook() -> bytes:
    return build_xlsx(
        {
            "Q3": [["Region", "Q3", "Q4"], ["EMEA", 1200, 1450], ["APAC", 980, None]],
            "Notes": [["owner"], ["ops@example.test"]],
        }
    )


@pytest.fixture
def formula_workbook() -> bytes:
    return build_xlsx({"Q3": [["a", "b"], [1, "=A2*2"]]})


@pytest.fixture
def literal_equals_workbook() -> bytes:
    """Text a user really types that HAPPENS to start with `=`.

    `=> see notes` is prose, not arithmetic. A spreadsheet stores it as an
    inline string (`data_type == "s"`), and every value-only guarantee this
    module makes applies to it exactly as to any other text.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q3"
    ws["A1"] = "note"
    ws["A2"] = "=> see notes"
    # What a spreadsheet application writes for typed text: openpyxl infers a
    # FORMULA from the leading `=` on assignment, so the type is set back to
    # a string the way the file on disk would have it.
    ws["A2"].data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def array_formula_workbook() -> bytes:
    """A formula openpyxl does NOT represent as a string.

    An array formula's `cell.value` is an `ArrayFormula` object, so a
    leading-`=` string test does not see it at all — while `data_type` is
    still `"f"`. This is the dangerous direction: a workbook that slips past
    the guard is edited, and saving blanks every cached value in it.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = 2
    ws["A2"] = 3
    ws["B1"] = ArrayFormula("B1:B1", "=SUM(A1:A2)")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def unicode_sheet_workbook() -> bytes:
    """Spaces, an em dash and non-ASCII — legal in Excel and awkward in a URL.

    NOT a slash: Excel forbids `\\ / ? * : [ ]` in sheet titles and openpyxl
    enforces it, so a slash-bearing sheet name cannot exist in a real
    workbook. Spaces and non-ASCII are the real encoding hazard.
    """
    return build_xlsx({"予算 — Q3 plan": [["名前", "値"], ["東京", 42]]})


@pytest.fixture
def dated_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dates"
    ws["A1"] = "when"
    ws["A2"] = dt.date(2026, 8, 22)
    ws["A3"] = dt.datetime(2026, 8, 22, 14, 30, 0)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
