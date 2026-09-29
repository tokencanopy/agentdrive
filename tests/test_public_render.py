"""The render layer turns artifact bytes into safe HTML.

Artifact content is UNTRUSTED — an agent wrote it. Raw HTML must never pass
through, or a shared artifact becomes stored XSS on our own origin.
"""

import pytest

from agentdrive.core.kinds import kind_for
from agentdrive.rendering.render import MAX_RENDER_BYTES, render_body


def test_markdown_renders_and_extracts_a_title():
    out = render_body(b"# Quarterly Report\n\nBody text.\n", "text/markdown", "q.md")
    assert out.mode == "markdown"
    assert "<h1>Quarterly Report</h1>" in out.html
    assert out.title == "Quarterly Report"


def test_raw_html_in_markdown_is_escaped_not_executed():
    out = render_body(b"Hello <script>alert(1)</script>\n", "text/markdown", "x.md")
    assert "<script>" not in out.html
    assert "&lt;script&gt;" in out.html


def test_code_is_highlighted():
    out = render_body(b"def f():\n    return 1\n", "text/x-python", "f.py")
    assert out.mode == "code"
    assert 'class="highlight"' in out.html


def test_image_renders_an_img_tag_pointing_at_the_content_route():
    out = render_body(b"", "image/png", "diagram.png")
    assert out.mode == "image"
    assert "<img" in out.html


def test_unknown_binary_falls_back_to_a_download_card():
    out = render_body(b"\x00\x01", "application/octet-stream", "blob.bin")
    assert out.mode == "download"
    assert "blob.bin" in out.html


def test_undecodable_bytes_do_not_raise():
    out = render_body(b"\xff\xfe\x00bad", "text/markdown", "b.md")
    assert out.mode == "download"


def test_oversized_text_falls_back_to_a_download_card():
    out = render_body(b"a" * (MAX_RENDER_BYTES + 1), "text/markdown", "big.md")
    assert out.mode == "download"


def test_markdown_without_a_heading_has_no_title():
    assert render_body(b"just a paragraph\n", "text/markdown", "p.md").title is None


@pytest.mark.parametrize(
    ("content_type", "name", "expected"),
    [
        # An uppercase extension is still markdown. Uploaders routinely send
        # text/plain for everything, so the filename has to carry the answer.
        ("text/plain", "README.MD", "markdown"),
        ("application/yaml", "compose.yaml", "data"),
        # Delimited data renders as a table, and the FILENAME is enough to
        # decide that — the classifier already called this a dataset on the
        # same evidence, so the body has to agree with the chip.
        ("application/octet-stream", "rows.csv", "table"),
        # content_type beats the filename: an image named .csv is an image.
        ("image/png", "chart.csv", "image"),
        ("video/mp4", "clip.mp4", "video"),
    ],
)
def test_mode_follows_the_shared_kind_classifier(content_type, name, expected):
    assert render_body(b"# t\n\nx", content_type, name).mode == expected


def test_mode_never_contradicts_the_kind_chip():
    """The body and the chip on the same page must agree about the artifact.

    Both derive from `kind_for`; this pins the mapping between them so a
    future edit to either cannot drift them apart silently.
    """
    mapping = {
        "md": "markdown",
        "image": "image",
        "code": "code",
        "dataset": "table",
        "video": "video",
    }
    for content_type, name in [
        ("text/markdown", "a.md"),
        ("image/png", "a.png"),
        ("application/json", "a.json"),
        ("text/csv", "a.csv"),
        ("video/mp4", "a.mp4"),
    ]:
        kind = kind_for(content_type, name)
        out = render_body(b"x", content_type, name)
        assert out.mode == mapping.get(kind, "download"), (content_type, kind)


def test_script_scheme_links_do_not_become_anchors():
    out = render_body(b"[click](javascript:alert(1))\n", "text/markdown", "l.md")
    assert "<a" not in out.html
    assert "javascript:" not in out.html.lower() or "href" not in out.html.lower()


def test_fenced_code_in_markdown_is_highlighted():
    """A fenced block must look the same as the equivalent code artifact.

    markdown-it's default is a bare `<pre><code class="language-x">`, which
    none of the stylesheet's token rules match — they are all scoped to
    `.highlight`. The same Python would render coloured as a `.py` artifact
    and monochrome inside a report.
    """
    out = render_body(b"```python\ndef f():\n    return 1\n```\n", "text/markdown", "a.md")
    assert out.mode == "markdown"
    assert 'class="highlight"' in out.html
    assert 'class="k"' in out.html  # `def` got a keyword token


def test_an_unknown_fence_language_still_escapes():
    """The fallback path must stay safe, not just unstyled.

    Returning "" hands the block back to markdown-it's own escaping. If that
    ever became "return the code raw", a fence labelled with a nonsense
    language would be an injection point straight past `html=False`.
    """
    out = render_body(
        b"```notalanguage\n<script>alert(1)</script>\n```\n", "text/markdown", "a.md"
    )
    assert "<script>" not in out.html
    assert "&lt;script&gt;" in out.html


def test_an_unlabelled_fence_still_escapes():
    out = render_body(b"```\n<img src=x onerror=alert(1)>\n```\n", "text/markdown", "a.md")
    assert "<img" not in out.html
    assert "&lt;img" in out.html


def test_a_fence_is_not_nested_inside_a_second_code_block():
    """markdown-it adopts a highlighter's output only if it starts with `<pre`.

    Pygments' own `cssclass=` wrapper is a `<div>`, which markdown-it would
    then nest inside its own `<pre><code>` — invalid markup, and visually a
    code block drawn inside another code block.
    """
    out = render_body(b"```python\ndef f():\n    pass\n```\n", "text/markdown", "a.md")
    from selectolax.parser import HTMLParser

    rendered = HTMLParser(out.html).css_first('[data-view="rendered"].doc-view').html
    assert "<pre><code" not in rendered
    assert rendered.count("<pre") == 1
    assert 'class="highlight"' in rendered


def test_a_pdf_renders_through_pdfjs():
    """The legacy engine, reused — not the browser's native `<embed>`.

    Legacy rejected `<embed>` deliberately: pdf.js gives selectable text,
    find-in-page and clickable links, which an opaque plugin box does not.
    That reasoning is why this reuses `static/pdfview.js` rather than
    reimplementing, and the element ids below are the ones its attach*
    helpers query — change them and the viewer silently stops wiring up.
    """
    out = render_body(b"%PDF-1.4 ...", "application/pdf", "report.pdf")

    assert out.mode == "pdf"
    assert "<embed" not in out.html
    assert 'id="pdf-doc"' in out.html
    assert 'data-pdf-url="content"' in out.html
    for el in ("pv-container", "pv-viewer", "pv-page-input", "pv-zoom-pct",
               "pv-find-input"):
        assert f'id="{el}"' in out.html, el


def test_a_pdf_page_degrades_without_javascript():
    """The CSP permits our script, but a reader may still have JS off.

    Without this they would get a blank frame and no way to reach the bytes.
    """
    out = render_body(b"%PDF-1.4 ...", "application/pdf", "report.pdf")
    assert "<noscript>" in out.html
    assert 'href="content" download' in out.html


def test_an_oversized_pdf_still_renders():
    """The size cap guards work WE do; pdf.js streams the document itself."""
    out = render_body(b"", "application/pdf", "big.pdf", size_bytes=50 * 1024 * 1024)
    assert out.mode == "pdf"


def test_the_print_button_is_drawn_only_when_a_surface_asked_for_it():
    """An unwired print button is a click that does nothing at all.

    `pdfview.js`'s `attachToolbar` takes every element as optional, so a
    surface that draws `pv-print` without wiring it gets a control that
    fails silently rather than a missing one. Only a caller that has wired
    the button may draw it, so the default is off.
    """
    off = render_body(b"%PDF-1.4 ...", "application/pdf", "report.pdf")
    assert 'id="pv-print"' not in off.html

    on = render_body(
        b"%PDF-1.4 ...", "application/pdf", "report.pdf", print_button=True
    )
    assert 'id="pv-print"' in on.html
    assert 'aria-label="Print"' in on.html


def test_print_is_a_pdf_only_control():
    """`print_button` names a PDF toolbar control; nothing else grows one."""
    for data, content_type, name in (
        (b"# hi\n", "text/markdown", "note.md"),
        (b"x,y\n1,2\n", "text/csv", "rows.csv"),
        (b"\x00\x01", "application/octet-stream", "blob.bin"),
    ):
        out = render_body(data, content_type, name, print_button=True)
        assert 'id="pv-print"' not in out.html, name


# ── delimited data renders as a table ────────────────────────────────────────


def test_csv_renders_as_a_table_with_a_header_row():
    out = render_body(b"a,b\n1,2\n3,4\n", "text/csv", "rows.csv")
    assert out.mode == "table"
    assert "<th>a</th><th>b</th>" in out.html
    assert "<td>1</td><td>2</td>" in out.html


def test_tsv_splits_on_tabs_not_commas():
    out = render_body(b"a\tb\n1,5\t2\n", "text/tab-separated-values", "rows.tsv")
    assert out.mode == "table"
    # The comma inside the first cell must stay inside it.
    assert "<td>1,5</td><td>2</td>" in out.html


def test_table_cells_are_escaped():
    out = render_body(b"a\n<script>x</script>\n", "text/csv", "x.csv")
    assert "<script>" not in out.html
    assert "&lt;script&gt;" in out.html


def test_table_is_row_capped_and_says_so():
    rows = b"h\n" + b"".join(b"%d\n" % i for i in range(600))
    out = render_body(rows, "text/csv", "big.csv")
    assert out.mode == "table"
    assert out.html.count("<tr>") == 501  # header + 500 data rows
    assert "500 of 600 rows" in out.html


def test_ragged_rows_are_padded_to_a_rectangle():
    """A short row must not shear the grid — it renders as empty cells."""
    out = render_body(b"a,b,c\n1\n", "text/csv", "ragged.csv")
    assert "<td>1</td><td></td><td></td>" in out.html


def test_unparseable_delimited_data_falls_back_to_text():
    """A NUL inside the field is a csv.Error; the text path is still useful."""
    out = render_body(b"a,b\n\x00\n", "text/csv", "bad.csv")
    assert out.mode in {"code", "table"}


# ── video ────────────────────────────────────────────────────────────────────


def test_small_video_plays_inline():
    out = render_body(b"\x00\x00\x00\x18ftyp", "video/mp4", "clip.mp4")
    assert out.mode == "video"
    assert 'class="artifact-video"' in out.html
    assert "controls" in out.html


def test_oversized_video_falls_back_to_the_download_card():
    """The embedded viewer fetches the whole blob, so the cap is a memory
    ceiling on the reader's tab, not a storage limit."""
    from agentdrive.rendering.render import MAX_INLINE_VIDEO_BYTES

    out = render_body(b"", "video/mp4", "big.mp4", size_bytes=MAX_INLINE_VIDEO_BYTES + 1)
    assert out.mode == "download"


# ── empty content ────────────────────────────────────────────────────────────


def test_empty_file_says_it_is_empty():
    out = render_body(b"", "text/plain", "nothing.txt")
    assert out.mode == "empty"
    assert "empty" in out.html.lower()


def test_whitespace_only_file_is_empty_too():
    out = render_body(b"   \n\t\n", "text/plain", "blank.txt")
    assert out.mode == "empty"


def test_empty_card_offers_no_download_link():
    """Nothing to download — offering one invites the reader to prove it.
    (The card reuses `.download-card` for its layout, so this asserts on the
    link, not on the class name.)"""
    out = render_body(b"", "text/markdown", "nothing.md")
    assert "<a" not in out.html
    assert "href" not in out.html


# ── remote images are refused visibly ────────────────────────────────────────


def test_remote_markdown_image_becomes_a_labelled_placeholder():
    out = render_body(
        b"![a chart](https://tracker.invalid/p.gif)\n", "text/markdown", "m.md"
    )
    assert "<img" not in out.html
    assert "blocked-remote" in out.html
    assert "a chart" in out.html


def test_local_markdown_image_is_left_alone():
    out = render_body(b"![ok](content)\n", "text/markdown", "m.md")
    assert '<img src="content"' in out.html


def test_unresolvable_markdown_image_becomes_a_labelled_placeholder():
    """A sibling artifact by name is the shape agents write, and it 404'd.

    `chart.png` resolved against the shell URL, where nothing is mounted, so
    a report that shipped with its chart showed a broken-image icon. Naming a
    sibling needs an addressing scheme we do not have; saying so does not.
    """
    out = render_body(b"![a chart](chart.png)\n", "text/markdown", "m.md")
    assert "<img" not in out.html
    assert "image-unavailable" in out.html
    assert "a chart" in out.html


def test_root_relative_and_protocol_relative_images_are_also_refused():
    for src in (b"/img/x.png", b"//host.invalid/x.png", b"../up.png"):
        out = render_body(b"![x](" + src + b")\n", "text/markdown", "m.md")
        assert "<img" not in out.html, src
        assert "image-unavailable" in out.html, src


def test_data_uri_images_are_left_alone():
    """Carried in the document, fetched from nowhere — and `img-src` allows it."""
    out = render_body(
        b"![dot](data:image/gif;base64,R0lGODlhAQABAAAAACw=)\n",
        "text/markdown",
        "m.md",
    )
    assert '<img src="data:image/gif;base64,' in out.html


def test_unavailable_placeholder_escapes_the_alt_text():
    out = render_body(
        b'![<script>x</script>](sibling.png)\n', "text/markdown", "m.md"
    )
    assert "<script>" not in out.html


def test_table_alignment_is_a_class_not_a_refused_style_attribute():
    """`style-src 'self'` refuses the attribute markdown-it emits by default."""
    out = render_body(
        b"| item | cost |\n| :--- | ---: |\n| a | 12 |\n", "text/markdown", "t.md"
    )
    assert "style=" not in out.html
    assert '<th class="ta-right">cost</th>' in out.html
    assert '<td class="ta-right">12</td>' in out.html
    assert '<td class="ta-left">a</td>' in out.html


def test_unaligned_table_cells_carry_no_alignment_class():
    out = render_body(b"| a |\n| --- |\n| 1 |\n", "text/markdown", "t.md")
    from selectolax.parser import HTMLParser

    rendered = HTMLParser(out.html).css_first('[data-view="rendered"].doc-view').html
    assert all("ta-" not in cell.attributes.get("class", "")
               for cell in HTMLParser(rendered).css("th, td"))


# ── GFM: the shapes every model emits ────────────────────────────────────────


def test_strikethrough_renders_rather_than_showing_tildes():
    out = render_body(b"a ~~gone~~ word\n", "text/markdown", "m.md")
    from selectolax.parser import HTMLParser

    rendered = HTMLParser(out.html).css_first('[data-view="rendered"].doc-view').html
    assert "<s>gone</s>" in rendered
    assert "~~" not in rendered


def test_task_lists_render_as_disabled_checkboxes():
    out = render_body(
        b"- [x] shipped\n- [ ] pending\n", "text/markdown", "m.md"
    )
    from selectolax.parser import HTMLParser

    rendered = HTMLParser(out.html).css_first('[data-view="rendered"].doc-view').html
    assert rendered.count("type=\"checkbox\"") == 2
    # Never interactive: the document is a read-only render of someone else's
    # file, and a checkbox that accepts a click promises a write we do not do.
    assert rendered.count('disabled="disabled"') == 2
    assert 'checked="checked"' in rendered
    assert "[x]" not in rendered


def test_task_list_markup_carries_no_author_html():
    """`html=False` is the whole security story; the plugin must not widen it."""
    out = render_body(
        b"- [x] <script>alert(1)</script>\n", "text/markdown", "m.md"
    )
    assert "<script>" not in out.html


def test_blocked_placeholder_escapes_the_alt_text():
    out = render_body(
        b'![<script>x</script>](http://elsewhere.invalid/a.png)\n',
        "text/markdown",
        "m.md",
    )
    assert "<script>" not in out.html


# ── markdown tables ──────────────────────────────────────────────────────────


def test_pipe_tables_render_as_tables():
    """CommonMark has no table syntax, so a GFM pipe table used to render as a
    paragraph of pipes — the shape agents write most, since every model emits
    GFM."""
    out = render_body(
        b"| a | b |\n| - | - |\n| 1 | 2 |\n", "text/markdown", "t.md"
    )
    assert out.mode == "markdown"
    assert "<table>" in out.html
    assert "<th>a</th>" in out.html
    assert "<td>1</td>" in out.html


def test_bare_urls_still_do_not_become_links():
    """Enabling `table` must not drag in linkify: an unlinked URL in untrusted
    text stays text, so the renderer never authors an anchor the document did
    not ask for."""
    out = render_body(b"see https://example.invalid/x for more\n", "text/markdown", "u.md")
    assert "<a" not in out.html


# ── workbooks ────────────────────────────────────────────────────────────
#
# Before the `sheet` mode every spreadsheet in the product answered with a
# download card: an xlsx is a ZIP archive, so it can only ever fail the UTF-8
# decode the text path starts with.


def _workbook(sheets: dict[str, list[list[object]]]) -> bytes:
    """A real xlsx, built in memory so no binary fixture lives in the tree."""
    import io as _io

    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_MODEL = {
    "Model": [["Region", "Q1", "Q2"], ["EMEA", 1200, 1310], ["AMER", 2400, 2510]],
    "Assumptions": [["Driver", "Value"], ["Growth", 0.08]],
}


def test_a_workbook_renders_as_a_sheet_not_a_download_card():
    out = render_body(_workbook(_MODEL), XLSX, "q3.xlsx")
    assert out.mode == "sheet"
    assert "EMEA" in out.html
    assert "1200" in out.html
    assert "download-card" not in out.html


def test_the_grid_carries_column_letters_and_row_numbers():
    # Without them a value has no address, and an address is how a reader
    # ties what they see here to the file they have open elsewhere.
    out = render_body(_workbook(_MODEL), XLSX, "q3.xlsx")
    assert ">A</th>" in out.html
    assert ">C</th>" in out.html
    assert 'class="workbook-rownum" scope="row">1</th>' in out.html


def test_every_sheet_is_rendered_with_a_tab_strip():
    out = render_body(_workbook(_MODEL), XLSX, "q3.xlsx")
    # Every sheet ships in the one document, so switching needs no request
    # and no script — which is what lets the console's viewer, whose shell
    # holds no credential after painting, switch tabs at all.
    assert 'id="sheet-0"' in out.html
    assert 'id="sheet-1"' in out.html
    assert "EMEA" in out.html
    assert "Growth" in out.html
    assert 'href="#sheet-1"' in out.html


def test_each_panel_marks_its_own_tab_current():
    # The strip is repeated per panel precisely so this is true with no
    # script: CSS cannot style an anchor from the element it points at.
    out = render_body(_workbook(_MODEL), XLSX, "q3.xlsx")
    assert out.html.count('aria-current="page"') == 2


def test_tabs_are_in_document_fragments_not_navigations():
    # `interceptLinks` in the console's viewer passes `#` hrefs through and
    # swallows everything else. A `?sheet=` tab would be a dead control.
    out = render_body(_workbook(_MODEL), XLSX, "q3.xlsx")
    assert "?sheet=" not in out.html
    assert 'href="#sheet-0"' in out.html


def test_sheet_names_are_escaped_wherever_they_appear():
    # A sheet name is agent-authored. Excel forbids `\ / ? * : [ ]` in a
    # title but not `&` or a quote, and both would break out of an attribute.
    # The href never carries the name at all — it is an index — so the tab
    # target cannot be injected into either.
    out = render_body(
        _workbook({"A&B \"x\"": [["v"]], "Other": [["w"]]}), XLSX, "q.xlsx"
    )
    assert "<script>" not in out.html
    assert "A&amp;B &quot;x&quot;" in out.html
    assert 'href="#sheet-0"' in out.html


def test_cell_values_are_escaped():
    out = render_body(
        _workbook({"S": [["<script>alert(1)</script>"]], "T": [["x"]]}),
        XLSX,
        "x.xlsx",
    )
    assert "<script>" not in out.html
    assert "&lt;script&gt;" in out.html


def test_truncation_is_stated_never_silent():
    rows = [[f"r{i}"] for i in range(600)]
    out = render_body(_workbook({"Big": rows}), XLSX, "big.xlsx")
    assert "500 of 600 rows" in out.html
    assert "download the file for all of it" in out.html


def test_wide_sheets_state_their_column_truncation_too():
    out = render_body(
        _workbook({"Wide": [[f"c{i}" for i in range(60)]]}), XLSX, "wide.xlsx"
    )
    assert "40 of 60 columns" in out.html


def test_a_workbook_we_cannot_open_still_offers_a_download():
    # Untrusted bytes: anything openpyxl refuses must degrade, not raise.
    out = render_body(b"not a zip at all", XLSX, "broken.xlsx")
    assert out.mode == "download"


def test_a_macro_workbook_keeps_the_download_card():
    # We decline to open macro containers at all, so there is nothing to draw.
    out = render_body(
        _workbook(_MODEL), "application/vnd.ms-excel.sheet.macroenabled.12", "m.xlsm"
    )
    assert out.mode == "download"


def test_an_oversized_workbook_is_never_parsed():
    out = render_body(b"", XLSX, "huge.xlsx", size_bytes=MAX_RENDER_BYTES + 1)
    assert out.mode == "download"


def test_a_legacy_xls_falls_back_to_a_download_card():
    """openpyxl cannot read BIFF, so `.xls` takes the workbook branch and
    lands on the fallback. The outcome is right; pin it so a future change
    to the suffix list cannot turn it into a stack trace."""
    out = render_body(b"\xd0\xcf\x11\xe0legacy binary", "application/vnd.ms-excel", "old.xls")
    assert out.mode == "download"


def test_the_viewer_grid_does_not_borrow_the_design_systems_class_names():
    """`.sheet-grid` belongs to the design system's HUNK table. AgentDrive is
    slated to adopt that stylesheet (AGENTS.md), and two meanings for one
    class in one document is a collision waiting to happen."""
    out = render_body(_workbook(_MODEL), XLSX, "q3.xlsx")
    for borrowed in ('class="sheet-grid"', "sheet-grid-rownum", 'class="sheet-tab'):
        assert borrowed not in out.html
    assert "workbook-grid" in out.html


def test_csv_still_renders_as_a_table_not_a_sheet():
    # The `table` mode is a contract with the stylesheet and both adapters;
    # folding csv into `sheet` is a separate, breaking change.
    out = render_body(b"a,b\n1,2\n", "text/csv", "d.csv")
    assert out.mode == "table"


def test_the_workbook_budget_bounds_every_sheet_together():
    """Per-sheet caps are not enough once every tab ships in one document.

    Six sheets at the 500-row cap would be 3,000 rows in one DOM; the
    workbook ceiling stops at 2,000 and each sheet states its own shortfall,
    so a reader is never told a sheet is complete when it was cut.
    """
    sheets = {f"S{i}": [[f"r{r}"] for r in range(500)] for i in range(6)}
    out = render_body(_workbook(sheets), XLSX, "many.xlsx")
    assert out.mode == "sheet"
    # 2,000 data rows plus one header row per sheet.
    assert out.html.count("workbook-rownum") == 2000
    assert "download the file for all of it" in out.html


def test_a_sheet_cut_short_by_the_budget_still_says_so():
    sheets = {"First": [[f"r{r}"] for r in range(1900)], "Second": [["a"], ["b"]]}
    out = render_body(_workbook(sheets), XLSX, "tight.xlsx")
    # The first sheet takes its own 500 cap and states it.
    assert "500 of 1900 rows" in out.html


# --- Word documents -----------------------------------------------------------

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# A 1×1 transparent PNG — the smallest real image a browser will decode.
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d4944415478da63f8ffff3f0300050001ff5f4a800000000049454e44ae426082"
)


def _docx(body: str, *, rels: str = "", media: dict[str, bytes] | None = None) -> bytes:
    """A real docx, built in memory so no binary fixture lives in the tree.

    `body` is the inside of `<w:body>`; `rels` extra `<Relationship>` rows for
    `word/_rels/document.xml.rels` (hyperlinks, pictures); `media` files under
    `word/`. Includes a styles part so "Heading 1" resolves to `<h1>`.
    """
    import io as _io
    import zipfile

    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="png" ContentType="image/png"/>'
            f'<Override PartName="/word/document.xml" ContentType="{DOCX}.main+xml"/>'
            '<Override PartName="/word/styles.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
            "</Types>",
        )
        zf.writestr(
            "_rels/.rels",
            f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{_REL_NS}">'
            f'<Relationship Id="rId1" Type="{_REL}/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        zf.writestr(
            "word/_rels/document.xml.rels",
            f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{_REL_NS}">'
            f'<Relationship Id="rIdStyles" Type="{_REL}/styles" Target="styles.xml"/>'
            f"{rels}</Relationships>",
        )
        zf.writestr(
            "word/styles.xml",
            f'<?xml version="1.0" encoding="UTF-8"?><w:styles {_W}>'
            '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>'
            "</w:styles>",
        )
        zf.writestr(
            "word/document.xml",
            f'<?xml version="1.0" encoding="UTF-8"?><w:document {_W} {_R} '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f"<w:body>{body}</w:body></w:document>",
        )
        for path, data in (media or {}).items():
            zf.writestr(f"word/{path}", data)
    return buf.getvalue()


def _p(text: str, *, style: str | None = None, bold: bool = False) -> str:
    from xml.sax.saxutils import escape as _xml

    text = _xml(text)
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    rpr = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f'<w:p>{ppr}<w:r>{rpr}<w:t xml:space="preserve">{text}</w:t></w:r></w:p>'


_REPORT = (
    _p("Quarterly Review", style="Heading1")
    + _p("Revenue grew.", bold=True)
    + "<w:tbl><w:tr><w:tc>" + _p("Region") + "</w:tc><w:tc>" + _p("Q1") + "</w:tc></w:tr>"
    + "<w:tr><w:tc>" + _p("EMEA") + "</w:tc><w:tc>" + _p("1200") + "</w:tc></w:tr></w:tbl>"
)


def test_a_word_document_renders_as_a_document_not_a_download_card():
    out = render_body(_docx(_REPORT), DOCX, "review.docx")
    assert out.mode == "document"
    assert "<h1>Quarterly Review</h1>" in out.html
    assert "<strong>Revenue grew.</strong>" in out.html
    assert "<table>" in out.html and "<td><p>EMEA</p></td>" in out.html
    assert 'class="download-card"' not in out.html


def test_a_word_document_is_titled_by_its_first_heading():
    out = render_body(_docx(_REPORT), DOCX, "review.docx")
    assert out.title == "Quarterly Review"
    assert render_body(_docx(_p("No heading here.")), DOCX, "n.docx").title is None


def test_a_word_document_has_no_source_pane_or_untrusted_page_strip():
    # There is no text form of a docx a reader could want, and the document
    # renders in our typography, not the author's — so no toggle, and no `page`.
    out = render_body(_docx(_REPORT), DOCX, "review.docx")
    assert 'data-view="source"' not in out.html
    assert 'class="doc-modes"' not in out.html
    assert out.html.startswith('<div class="doc-view" data-view="rendered">')


def test_a_word_document_is_recognised_by_suffix_under_a_generic_type():
    out = render_body(_docx(_REPORT), "application/octet-stream", "review.docx")
    assert out.mode == "document"


def test_markup_inside_a_word_document_is_text_not_html():
    hostile = _docx(_p("<script>alert(1)</script> & <img src=x onerror=1>"))
    out = render_body(hostile, DOCX, "hostile.docx")
    assert out.mode == "document"
    assert "<script" not in out.html and "<img" not in out.html
    assert "&lt;script&gt;alert(1)&lt;/script&gt; &amp; &lt;img src=x onerror=1&gt;" in out.html


def test_word_hyperlinks_keep_https_and_drop_javascript():
    body = (
        '<w:p><w:hyperlink r:id="rId7"><w:r><w:t>safe</w:t></w:r></w:hyperlink>'
        '<w:hyperlink r:id="rId8"><w:r><w:t>hostile</w:t></w:r></w:hyperlink></w:p>'
    )
    rels = "".join(
        f'<Relationship Id="{rid}" Type="{_REL}/hyperlink" Target="{target}" '
        'TargetMode="External"/>'
        for rid, target in (("rId7", "https://example.com/x"), ("rId8", "javascript:alert(1)"))
    )
    out = render_body(_docx(body, rels=rels), DOCX, "links.docx")
    assert '<a href="https://example.com/x">safe</a>' in out.html
    assert "javascript:" not in out.html
    assert "hostile" in out.html  # the text survives; only the target is dropped


def test_an_embedded_picture_renders_inline_as_a_data_uri():
    body = (
        "<w:p><w:r><w:drawing><wp:inline>"
        '<wp:docPr id="1" name="Picture 1" descr="Revenue chart"/>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic><pic:blipFill><a:blip r:embed="rId9"/></pic:blipFill></pic:pic>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
    )
    rels = f'<Relationship Id="rId9" Type="{_REL}/image" Target="media/image1.png"/>'
    out = render_body(_docx(body, rels=rels, media={"media/image1.png": _PNG}), DOCX, "pic.docx")
    assert out.mode == "document"
    assert '<img alt="Revenue chart" src="data:image/png;base64,' in out.html


def test_an_empty_word_document_says_so():
    out = render_body(_docx(_p("")), DOCX, "blank.docx")
    assert out.mode == "empty"


def test_a_word_document_we_cannot_open_still_offers_a_download():
    # Untrusted bytes: anything mammoth refuses must degrade, not raise.
    assert render_body(b"not a zip at all", DOCX, "broken.docx").mode == "download"
    # A valid zip that is not a Word package — no document part inside.
    import io as _io
    import zipfile

    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "not a document")
    assert render_body(buf.getvalue(), DOCX, "hollow.docx").mode == "download"


def test_macro_and_legacy_word_files_keep_the_download_card():
    # `.docm` is refused like `.xlsm`; `.doc` is BIFF, which mammoth cannot read.
    docm = "application/vnd.ms-word.document.macroenabled.12"
    assert render_body(_docx(_REPORT), docm, "m.docm").mode == "download"
    legacy = render_body(b"\xd0\xcf\x11\xe0legacy binary", "application/msword", "old.doc")
    assert legacy.mode == "download"


def test_an_oversized_word_document_is_never_parsed():
    out = render_body(b"", DOCX, "huge.docx", size_bytes=MAX_RENDER_BYTES + 1)
    assert out.mode == "download"


def test_content_the_converter_drops_is_stated_not_silent():
    # OMML equations are the common case: mammoth ignores them, so a report
    # whose numbers live in its formulas would otherwise read as complete.
    body = _p("The model is:") + (
        '<w:p><m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:oMath><m:r><m:t>E=mc2</m:t></m:r></m:oMath></m:oMathPara></w:p>"
    )
    out = render_body(_docx(body), DOCX, "physics.docx")
    assert out.mode == "document"
    assert '<p class="doc-note">Not shown in this preview: equations.' in out.html
    # And nothing is said when nothing was dropped.
    assert "doc-note" not in render_body(_docx(_REPORT), DOCX, "review.docx").html
# --- Parquet ------------------------------------------------------------------

PARQUET = "application/vnd.apache.parquet"


def _parquet(columns: dict[str, list[object]], **kwargs) -> bytes:
    """A real parquet file, built in memory so no binary fixture lives in the tree."""
    import io as _io

    import pyarrow as pa
    import pyarrow.parquet as pq

    buf = _io.BytesIO()
    pq.write_table(pa.table(columns), buf, **kwargs)
    return buf.getvalue()


_ROWS = {"region": ["EMEA", "AMER"], "q1": [1200, 2400], "growth": [0.08, None]}


def test_parquet_renders_as_the_same_table_a_csv_does():
    out = render_body(_parquet(_ROWS), PARQUET, "rows.parquet")
    assert out.mode == "table"
    assert '<table class="data-table">' in out.html
    assert '<th title="string">region</th>' in out.html
    assert '<th title="int64">q1</th>' in out.html
    assert "<td>EMEA</td><td>1200</td><td>0.08</td>" in out.html
    assert "<td>AMER</td><td>2400</td><td></td>" in out.html  # null is empty, not "None"
    assert 'class="download-card"' not in out.html


def test_parquet_has_no_raw_text_pane():
    # A csv's "Raw text" is its own bytes; parquet has no text a reader could want.
    out = render_body(_parquet(_ROWS), PARQUET, "rows.parquet")
    assert 'data-view="source"' not in out.html
    assert 'class="doc-modes"' not in out.html


def test_parquet_is_recognised_by_suffix_under_a_generic_type():
    assert render_body(_parquet(_ROWS), "application/octet-stream", "rows.parquet").mode == "table"
    assert render_body(_parquet(_ROWS), "application/x-parquet", "rows").mode == "table"


def test_parquet_cells_are_escaped():
    out = render_body(_parquet({"c": ["<img src=x onerror=1>"]}), PARQUET, "x.parquet")
    assert "<img" not in out.html
    assert "&lt;img src=x onerror=1&gt;" in out.html


def test_parquet_is_row_capped_and_states_the_true_total_from_the_footer():
    n = 1200
    out = render_body(_parquet({"i": list(range(n))}, row_group_size=n), PARQUET, "big.parquet")
    assert out.mode == "table"
    assert out.html.count("<tr>") == 1 + 500
    assert "Showing 500 of 1200 rows" in out.html


def test_parquet_reads_only_the_first_row_group_but_counts_them_all():
    out = render_body(_parquet({"i": list(range(300))}, row_group_size=100), PARQUET, "g.parquet")
    assert out.html.count("<tr>") == 1 + 100
    assert "Showing 100 of 300 rows" in out.html


def test_wide_parquet_states_its_column_truncation_too():
    out = render_body(_parquet({f"c{i}": [i] for i in range(45)}), PARQUET, "wide.parquet")
    assert "Showing 40 of 45 columns" in out.html
    assert "<th title=" in out.html and "c44" not in out.html


def test_parquet_typed_cells_render_as_a_reader_would_write_them():
    import datetime as dt
    import decimal

    out = render_body(
        _parquet({
            "flag": [True],
            "when": [dt.datetime(2026, 1, 2, 3, 4, 5)],
            "day": [dt.date(2026, 1, 2)],
            "blob": [b"\x00\x01\x02"],
            "tags": [["a", "b"]],
            "money": [decimal.Decimal("12.50")],
        }),
        PARQUET,
        "typed.parquet",
    )
    assert "<td>true</td>" in out.html
    assert "<td>2026-01-02T03:04:05</td>" in out.html
    assert "<td>2026-01-02</td>" in out.html
    assert "<td>3 bytes</td>" in out.html
    assert '<td>[&quot;a&quot;, &quot;b&quot;]</td>' in out.html
    assert "<td>12.50</td>" in out.html


def test_a_long_parquet_cell_is_elided():
    out = render_body(_parquet({"text": ["x" * 5000]}), PARQUET, "long.parquet")
    assert "x" * 500 + "…" in out.html
    assert "x" * 501 not in out.html


def test_an_empty_parquet_renders_its_header():
    out = render_body(_parquet({"region": pa_empty_strings()}), PARQUET, "empty.parquet")
    assert out.mode == "table"
    assert "<th title=" in out.html and out.html.count("<tr>") == 1


def pa_empty_strings():
    import pyarrow as pa

    return pa.array([], type=pa.string())


def test_parquet_we_cannot_open_still_offers_a_download():
    # Untrusted bytes: anything pyarrow refuses must degrade, not raise.
    assert render_body(b"PAR1 not really PAR1", PARQUET, "broken.parquet").mode == "download"
    assert render_body(_workbook(_MODEL), PARQUET, "actually.xlsx.parquet").mode == "download"


def test_an_oversized_parquet_is_never_parsed():
    out = render_body(b"", PARQUET, "huge.parquet", size_bytes=MAX_RENDER_BYTES + 1)
    assert out.mode == "download"


def test_a_row_group_claiming_more_than_the_guard_is_refused(monkeypatch):
    from agentdrive.rendering import render as render_module

    monkeypatch.setattr(render_module, "_PARQUET_MAX_ROW_GROUP_BYTES", 1)
    assert render_body(_parquet(_ROWS), PARQUET, "rows.parquet").mode == "download"
# --- Diagrams -----------------------------------------------------------------

MERMAID_DOC = b"# Plan\n\n```mermaid\ngraph TD; A-->B;\n```\n\nText after.\n"


def test_a_mermaid_fence_is_marked_for_the_client_not_highlighted():
    out = render_body(MERMAID_DOC, "text/markdown", "plan.md")
    assert out.mode == "markdown"
    assert out.diagrams is True
    # The source is escaped and carried in the document — it is the fallback
    # with script off and the text the engine reads with it on.
    assert (
        '<pre class="diagram-source" data-diagram="mermaid"><code>graph TD; A--&gt;B;\n'
        "</code></pre>"
    ) in out.html
    assert 'class="language-mermaid"' not in out.html


def test_the_fence_language_is_matched_case_insensitively():
    out = render_body(b"```Mermaid\nflowchart LR; a-->b\n```\n", "text/markdown", "d.md")
    assert out.diagrams is True


def test_documents_without_a_diagram_do_not_ask_for_the_engine():
    assert render_body(b"# Plain\n\ntext\n", "text/markdown", "p.md").diagrams is False
    assert render_body(b"```python\nx = 1\n```\n", "text/markdown", "c.md").diagrams is False
    assert render_body(b"x = 1\n", "text/x-python", "c.py").diagrams is False
    # The marker written as prose is escaped text, not a diagram.
    prose = render_body(b'Write data-diagram="mermaid" here.\n', "text/markdown", "n.md")
    assert prose.diagrams is False


def test_an_html_artifact_cannot_forge_a_diagram(monkeypatch):
    # Only the markdown fence can emit the marker: the sanitiser drops every
    # data-* attribute an author writes, so `page` mode never asks for the
    # engine and never hands author text to it.
    from agentdrive.rendering import render as render_module

    monkeypatch.setattr(render_module, "_static_html_enabled", lambda: True)
    forged = b'<pre class="diagram-source" data-diagram="mermaid"><code>graph TD</code></pre>'
    out = render_body(forged, "text/html", "forged.html")
    assert out.mode == "page"
    assert out.diagrams is False
    assert "data-diagram" not in out.html.split('data-view="source"')[0]


# --- JSON ---------------------------------------------------------------------


def test_json_renders_as_a_foldable_tree_with_its_source_a_click_away():
    doc = b'{"name": "run-7", "stats": {"ok": 12, "failed": 0}, "tags": ["a", "b"], "done": true}'
    out = render_body(doc, "application/json", "run.json")
    assert out.mode == "data"
    assert '<div class="json-tree">' in out.html
    assert '<span class="json-key">&quot;name&quot;</span>' in out.html
    assert '<span class="json-string">&quot;run-7&quot;</span>' in out.html
    assert '<span class="json-number">12</span>' in out.html
    assert '<span class="json-bool">true</span>' in out.html
    # Nested containers are native disclosures — fold without script.
    assert '<details class="json-node" open><summary>' in out.html
    assert '<span class="json-count">2 fields</span>' in out.html
    # The highlighted source is still there, behind the strip.
    assert 'data-view="rendered"' in out.html and 'data-view="source"' in out.html
    assert ">Tree<" in out.html and ">Source<" in out.html
    assert 'class="highlight"' in out.html


def test_json_records_render_as_the_shared_table():
    doc = b'[{"id": 1, "region": "EMEA"}, {"id": 2, "region": "AMER", "note": "late"}]'
    out = render_body(doc, "application/json", "rows.json")
    assert out.mode == "table"
    assert '<table class="data-table">' in out.html
    assert "<th>id</th><th>region</th><th>note</th>" in out.html
    assert "<td>1</td><td>EMEA</td><td></td>" in out.html
    assert ">Table<" in out.html and ">Source<" in out.html


def test_json_lines_of_records_render_as_the_shared_table():
    doc = b'{"id": 1, "ok": true}\n{"id": 2, "ok": false}\n\n'
    out = render_body(doc, "application/x-ndjson", "events.jsonl")
    assert out.mode == "table"
    assert "<td>1</td><td>true</td>" in out.html and "<td>2</td><td>false</td>" in out.html
    # The suffix alone is enough under a generic type.
    assert render_body(doc, "text/plain", "events.jsonl").mode == "table"
    assert render_body(b'{"a": 1}', "text/plain", "thing.json").mode == "data"


def test_json_records_with_nested_values_are_a_tree_not_a_table():
    doc = b'[{"id": 1, "tags": ["x"]}, {"id": 2, "tags": []}]'
    out = render_body(doc, "application/json", "nested.json")
    assert out.mode == "data"
    assert '<span class="json-count">2 items</span>' in out.html


def test_json_keeps_number_lexemes_key_order_and_duplicate_keys():
    doc = b'{"z": 1.0, "a": 1e3, "z": 2}'
    out = render_body(doc, "application/json", "dup.json")
    assert '<span class="json-number">1.0</span>' in out.html
    assert '<span class="json-number">1e3</span>' in out.html
    assert out.html.count('<span class="json-key">&quot;z&quot;</span>') == 2
    assert out.html.index("&quot;z&quot;") < out.html.index("&quot;a&quot;")


def test_json_strings_are_escaped_not_markup():
    doc = b'{"html": "<img src=x onerror=alert(1)>", "<b>key</b>": "v"}'
    out = render_body(doc, "application/json", "hostile.json")
    assert "<img" not in out.html and "<b>" not in out.html
    assert "&lt;img src=x onerror=alert(1)&gt;" in out.html


def test_json_tree_is_bounded_and_says_so():
    from agentdrive.rendering import data as jsondata

    doc = ("[" + ",".join(str(i) for i in range(20_000)) + "]").encode()
    out = render_body(doc, "application/json", "big.json")
    assert out.mode == "data"
    assert out.html.count('<div class="json-leaf">') == jsondata.MAX_ITEMS
    assert f"… {20_000 - jsondata.MAX_ITEMS:,} more items" in out.html
    assert '<p class="doc-note">Showing the first 500 items of long arrays' in out.html
    deep = ("[" * 40 + "1" + "]" * 40).encode()
    assert "16 levels of nesting" in render_body(deep, "application/json", "deep.json").html
    long = json_bytes({"s": "x" * 5000})
    out = render_body(long, "application/json", "long.json")
    rendered = out.html.split('<div class="doc-view" data-view="source"')[0]  # source keeps all
    assert "x" * 500 + "…" in rendered and "x" * 501 not in rendered
    assert "long strings" in rendered


def json_bytes(value) -> bytes:
    import json as _json

    return _json.dumps(value).encode()


def test_json_records_table_is_row_capped_with_the_true_total():
    doc = json_bytes([{"i": i} for i in range(1200)])
    out = render_body(doc, "application/json", "many.json")
    assert out.mode == "table"
    assert out.html.count("<tr>") == 1 + 500
    assert "Showing 500 of 1200 rows" in out.html


def test_invalid_json_still_renders_as_highlighted_source():
    out = render_body(b'{"unterminated": ', "application/json", "broken.json")
    assert out.mode == "code"
    assert 'class="highlight"' in out.html


def test_empty_json_containers_and_scalars_render():
    empty = render_body(b"{}", "application/json", "e.json")
    assert '<span class="json-brace">{}</span>' in empty.html
    assert render_body(b"[]", "application/json", "e.json").mode == "data"
    out = render_body(b'"just a string"', "application/json", "s.json")
    assert out.mode == "data" and "just a string" in out.html
# --- Oversized tabular files, previewed from a fraction ------------------------

from agentdrive.rendering.source import BytesSource  # noqa: E402


def _big_parquet(rows: int, groups: int) -> bytes:
    import io as _io

    import pyarrow as pa
    import pyarrow.parquet as pq

    buf = _io.BytesIO()
    table = pa.table({"i": list(range(rows)), "s": [f"row-{i}" for i in range(rows)]})
    pq.write_table(table, buf, row_group_size=max(1, rows // groups), compression="none")
    return buf.getvalue()


def test_an_oversized_parquet_previews_from_its_footer_and_one_row_group():
    data = _big_parquet(rows=300_000, groups=30)
    assert len(data) > MAX_RENDER_BYTES
    source = BytesSource(data)
    out = render_body(b"", PARQUET, "big.parquet", size_bytes=len(data), source=source)
    assert out.mode == "table"
    assert out.html.count("<tr>") == 1 + 500
    assert "Showing 500 of 300000 rows" in out.html  # the footer knows the true total
    # The point: a fraction of the file, not the file.
    assert source.fetched < len(data) // 10, (source.fetched, len(data))
    assert source.calls <= 6


def test_an_oversized_parquet_without_a_source_is_still_a_download():
    out = render_body(b"", PARQUET, "big.parquet", size_bytes=MAX_RENDER_BYTES + 1)
    assert out.mode == "download"


def test_an_oversized_csv_previews_from_its_first_megabyte():
    from agentdrive.rendering.render import HEAD_PREVIEW_BYTES

    lines = ["id,name"] + [f"{i},row-{i}" for i in range(400_000)]
    data = "\n".join(lines).encode()
    assert len(data) > MAX_RENDER_BYTES
    source = BytesSource(data)
    out = render_body(b"", "text/csv", "big.csv", size_bytes=len(data), source=source)
    assert out.mode == "table"
    assert out.html.count("<tr>") == 1 + 500
    assert "<td>0</td><td>row-0</td>" in out.html
    # One head read, and the caption says what was read rather than guess a total.
    assert source.calls == 1 and source.fetched == HEAD_PREVIEW_BYTES
    assert "Showing 500 rows from the first 1.0 MB of a" in out.html
    assert "MB file" in out.html
    # No raw-text pane for a partial read: the bytes on hand are not the file.
    assert 'data-view="source"' not in out.html


def test_the_cut_line_of_a_csv_head_is_never_shown_as_a_row(monkeypatch):
    from agentdrive.rendering import render as render_module

    # A head that ends mid-row. The head size is shrunk so the cut lands inside
    # the third line; the file itself is far over the render cap.
    head = b"id,name\n1,alpha\n2,beta\n3,gam"
    source = BytesSource(head + b"ma\n" + b"4,delta\n" * 400_000)
    assert source.size > MAX_RENDER_BYTES
    monkeypatch.setattr(render_module, "HEAD_PREVIEW_BYTES", len(head))
    out = render_body(b"", "text/csv", "cut.csv", size_bytes=source.size, source=source)
    assert out.mode == "table"
    assert "<td>2</td><td>beta</td>" in out.html
    assert "gam" not in out.html  # the partial third row is dropped, not shown truncated


def test_a_large_preview_that_would_exceed_the_read_budget_is_a_download():
    from agentdrive.rendering import source as source_module

    data = _big_parquet(rows=300_000, groups=1)
    assert len(data) > MAX_RENDER_BYTES
    original = source_module.RANGED_READ_BUDGET
    try:
        source_module.RANGED_READ_BUDGET = 1024
        # RangedFile reads its default at construction time from the module.
        from agentdrive.rendering import render as render_module

        real = render_module.RangedFile
        render_module.RangedFile = lambda s: real(s, budget=1024)
        try:
            out = render_body(
                b"", PARQUET, "b.parquet", size_bytes=len(data), source=BytesSource(data)
            )
        finally:
            render_module.RangedFile = real
    finally:
        source_module.RANGED_READ_BUDGET = original
    assert out.mode == "download"


def test_oversized_non_tabular_formats_still_card_even_with_a_source():
    from agentdrive.rendering.render import previews_when_large

    assert previews_when_large("text/csv", "a.csv")
    assert previews_when_large("application/vnd.apache.parquet", "a.parquet")
    assert previews_when_large("text/markdown", "a.md")  # text previews its head
    assert not previews_when_large(XLSX, "a.xlsx")
    book = BytesSource(b"PK\x03\x04" + b"x" * 500_000)
    out = render_body(b"", XLSX, "big.xlsx", size_bytes=64 * 1024 * 1024, source=book)
    assert out.mode == "download" and book.calls == 0


# --- Slide decks --------------------------------------------------------------

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _deck(*, chart: bool = False, picture: bytes | None = None, slides: int = 1) -> bytes:
    """A real pptx, built in memory so no binary fixture lives in the tree."""
    import io as _io

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for n in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Quarterly Review" if n == 0 else f"Slide {n + 1} title"
        body = slide.placeholders[1].text_frame
        body.text = "Revenue grew"
        para = body.add_paragraph()
        para.text = "EMEA up 8% <b>"
        para.level = 1
        run = para.add_run()
        run.text = " details"
        run.hyperlink.address = "https://example.com/x"
        run2 = para.add_run()
        run2.text = " hostile"
        run2.hyperlink.address = "javascript:alert(1)"
        if n == 0:
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
            table.cell(0, 0).text = "Region"
            table.cell(0, 1).text = "Q1"
            table.cell(1, 0).text = "EMEA"
            table.cell(1, 1).text = "1200"
            slide.notes_slide.notes_text_frame.text = "Remember to mention churn."
            if picture is not None:
                slide.shapes.add_picture(_io.BytesIO(picture), Inches(5), Inches(1))
            if chart:
                from pptx.chart.data import CategoryChartData
                from pptx.enum.chart import XL_CHART_TYPE

                data = CategoryChartData()
                data.categories = ["A", "B"]
                data.add_series("S", (1, 2))
                slide.shapes.add_chart(
                    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4.5), Inches(4), Inches(2),
                    data,
                )
    buf = _io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_a_deck_renders_as_its_slides_content():
    out = render_body(_deck(picture=_PNG), PPTX, "review.pptx")
    assert out.mode == "document"
    assert out.title == "Quarterly Review"
    assert '<section class="slide" id="slide-1">' in out.html
    assert '<h2><span class="slide-number">1</span> Quarterly Review</h2>' in out.html
    # Body text becomes a nested list from paragraph levels.
    assert "<ul><li>Revenue grew</li><ul><li>EMEA up 8% &lt;b&gt;" in out.html
    # Hyperlinks: https kept, javascript: dropped with its text kept.
    assert '<a href="https://example.com/x" rel="noopener"> details</a>' in out.html
    assert "javascript:" not in out.html and " hostile" in out.html
    # Tables are the shared grid; notes fold behind a disclosure; pictures inline.
    assert '<thead><tr><th>Region</th><th>Q1</th></tr></thead>' in out.html
    assert "<tr><td>EMEA</td><td>1200</td></tr>" in out.html
    assert "<summary>Speaker notes</summary><p>Remember to mention churn.</p>" in out.html
    assert '<img alt="Picture' in out.html and 'src="data:image/png;base64,' in out.html
    assert "doc-note" not in out.html


def test_a_deck_names_what_it_cannot_show():
    out = render_body(_deck(chart=True), PPTX, "charts.pptx")
    assert out.mode == "document"
    assert '<p class="doc-note">Not shown in this preview: charts.' in out.html


def test_a_deck_is_recognised_by_suffix_under_a_generic_type():
    assert render_body(_deck(), "application/octet-stream", "review.pptx").mode == "document"


def test_macro_and_legacy_decks_keep_the_download_card():
    assert render_body(_deck(), "application/vnd.ms-powerpoint.presentation.macroenabled.12",
                       "m.pptm").mode == "download"
    assert render_body(b"\xd0\xcf\x11\xe0legacy", "application/vnd.ms-powerpoint",
                       "old.ppt").mode == "download"


def test_a_deck_we_cannot_open_still_offers_a_download():
    assert render_body(b"not a zip", PPTX, "broken.pptx").mode == "download"
    assert render_body(_workbook(_MODEL), PPTX, "hollow.pptx").mode == "download"


def test_a_deck_over_the_text_cap_renders_whole_up_to_its_own_ceiling():
    from agentdrive.rendering import render as render_module
    from agentdrive.rendering.render import SLIDES_MAX_BYTES

    data = _deck(slides=3)
    # Under the deck ceiling the routes fetch the whole file (see
    # `preview_input`), so the renderer sees real bytes with a size past the
    # text cap — and renders them.
    out = render_body(data, PPTX, "big.pptx", size_bytes=MAX_RENDER_BYTES + 1)
    assert out.mode == "document"
    assert out.html.count('<section class="slide"') == 3
    # Past the deck ceiling it is a card without a single read, source or not.

    class Huge:
        size = SLIDES_MAX_BYTES + 1
        calls = 0

        def read_range(self, start, end):
            self.calls += 1
            return b""

    h = Huge()
    out = render_body(b"", PPTX, "huge.pptx", size_bytes=h.size, source=h)
    assert out.mode == "download" and h.calls == 0
    assert not render_module.previews_when_large(PPTX, "x.pptx")


def test_deck_pictures_are_bounded_and_the_note_counts_them(monkeypatch):
    from agentdrive.rendering import slides as slides_module

    monkeypatch.setattr(slides_module, "IMAGE_MAX_BYTES", 10)
    out = render_body(_deck(picture=_PNG), PPTX, "pics.pptx")
    assert "data:image/png" not in out.html
    assert 'aria-label="Image too large to show inline"' in out.html
    assert "1 picture too large to show inline" in out.html


def test_deck_slide_count_is_bounded_and_stated(monkeypatch):
    from agentdrive.rendering import slides as slides_module

    monkeypatch.setattr(slides_module, "MAX_SLIDES", 2)
    out = render_body(_deck(slides=3), PPTX, "long.pptx")
    assert out.html.count('<section class="slide"') == 2
    assert "the first 2 of 3 slides" in out.html


# --- YAML and XML -------------------------------------------------------------


def test_yaml_renders_as_the_same_foldable_tree_as_json():
    doc = (
        b"name: run-7\nversion: 1.0\nport: 0x1F\nstats:\n  ok: 12\n  failed: 0\n"
        b"tags: [a, b]\ndone: true\nwhen: 2026-01-02\nnothing: ~\n"
    )
    out = render_body(doc, "application/yaml", "run.yaml")
    assert out.mode == "data"
    assert '<div class="json-tree">' in out.html
    assert '<span class="json-key">&quot;name&quot;</span>' in out.html
    assert '<span class="json-string">&quot;run-7&quot;</span>' in out.html
    # Numbers are the lexeme the author wrote, not PyYAML's float or int.
    assert '<span class="json-number">1.0</span>' in out.html
    assert '<span class="json-number">0x1F</span>' in out.html
    assert '<span class="json-bool">true</span>' in out.html
    assert '<span class="json-null">null</span>' in out.html
    assert "2026-01-02" in out.html
    assert '<span class="json-count">2 fields</span>' in out.html
    assert ">Tree<" in out.html and ">Source<" in out.html and 'class="highlight"' in out.html


def test_yaml_records_render_as_the_shared_table():
    doc = b"- id: 1\n  region: EMEA\n- id: 2\n  region: AMER\n  note: late\n"
    out = render_body(doc, "text/yaml", "rows.yml")
    assert out.mode == "table"
    assert "<th>id</th><th>region</th><th>note</th>" in out.html
    assert "<td>1</td><td>EMEA</td><td></td>" in out.html
    # Suffix alone is enough under a generic type.
    assert render_body(doc, "text/plain", "rows.yml").mode == "table"


def test_multi_document_yaml_is_a_list_of_documents():
    out = render_body(b"a:\n  x: 1\n---\nb:\n  y: 2\n", "application/yaml", "multi.yaml")
    assert out.mode == "data"
    assert '<span class="json-count">2 items</span>' in out.html
    # Flat documents are records, and records are a table — the same rule JSON has.
    assert render_body(b"a: 1\n---\na: 2\n", "application/yaml", "flat.yaml").mode == "table"


def test_yaml_refuses_arbitrary_tags_and_keeps_the_source():
    # `!!python/object` would construct objects under the full loader; the
    # safe loader raises and the file stays highlighted source.
    out = render_body(b"!!python/object/apply:os.system [echo]\n", "application/yaml", "evil.yaml")
    assert out.mode == "code"
    assert 'class="highlight"' in out.html
    assert render_body(b"key: [unclosed", "application/yaml", "broken.yaml").mode == "code"


def test_yaml_strings_and_keys_are_escaped():
    out = render_body(b'"<b>k</b>": "<img src=x onerror=alert(1)>"\n', "application/yaml", "h.yaml")
    assert "<img" not in out.html and "<b>" not in out.html
    assert "&lt;img src=x onerror=alert(1)&gt;" in out.html


def test_xml_renders_as_an_element_tree():
    doc = (
        b'<?xml version="1.0"?>\n<!-- a comment -->\n'
        b'<report xmlns:x="urn:example" id="r7">\n  <title>Quarterly</title>\n'
        b'  <x:region code="EMEA">up 8%</x:region>\n  <empty/>\n'
        b"  <items><item>1</item><item>2</item></items>\n</report>\n"
    )
    out = render_body(doc, "application/xml", "report.xml")
    assert out.mode == "data"
    assert '<div class="json-tree xml-tree">' in out.html
    assert '<span class="xml-tag">&lt;report</span>' in out.html
    assert '<span class="xml-attr">id</span>' in out.html
    assert '<span class="json-string">&quot;r7&quot;</span>' in out.html
    # Text-only elements are one-line leaves; empty ones self-close.
    assert "&lt;title</span><span class=\"xml-tag\">&gt;</span>" in out.html
    assert '<span class="xml-text">Quarterly</span>' in out.html
    assert "&lt;empty</span><span class=\"xml-tag\">/&gt;</span>" in out.html
    # Namespaced names show the local name with the namespace on hover.
    assert '<span class="xml-tag" title="urn:example">&lt;region</span>' in out.html
    assert '<span class="json-count">2 elements</span>' in out.html
    rendered = out.html.split('<div class="doc-view" data-view="source"')[0]
    assert "a comment" not in rendered  # comments are not the document's data
    assert ">Tree<" in out.html and 'class="highlight"' in out.html


def test_xml_is_recognised_by_generic_plus_xml_types_and_suffixes():
    doc = b"<feed><entry>x</entry></feed>"
    assert render_body(doc, "application/atom+xml", "feed").mode == "data"
    assert render_body(doc, "text/xml", "f").mode == "data"
    assert render_body(doc, "text/plain", "f.xml").mode == "data"
    assert render_body(doc, "application/octet-stream", "f.plist").mode == "data"
    # SVG stays an image: the type wins before this dispatch is reached.
    assert render_body(b"<svg/>", "image/svg+xml", "a.svg").mode == "image"


def test_xml_entity_bombs_and_external_entities_are_refused():
    bomb = (
        b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
        b'<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
        b'<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]>'
        b"<lolz>&lol3;</lolz>"
    )
    out = render_body(bomb, "application/xml", "bomb.xml")
    assert out.mode == "code"  # refused by defusedxml; the source is shown instead
    external = (
        b'<?xml version="1.0"?><!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
        b"<foo>&xxe;</foo>"
    )
    assert render_body(external, "application/xml", "xxe.xml").mode == "code"
    assert render_body(b"<unclosed>", "application/xml", "broken.xml").mode == "code"


def test_xml_text_and_attributes_are_escaped():
    doc = b'<a href="javascript:alert(1)" onmouseover="x">&lt;script&gt;alert(1)&lt;/script&gt;</a>'
    out = render_body(doc, "application/xml", "hostile.xml")
    assert out.mode == "data"
    rendered = out.html.split('<div class="doc-view" data-view="source"')[0]
    assert "<script" not in rendered and 'href="javascript' not in rendered
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in rendered
    # Attributes are text in a span, never attributes on our elements.
    assert 'onmouseover="x"' not in rendered


def test_xml_tree_is_bounded_and_says_so():
    from agentdrive.rendering import data as jsondata

    doc = b"<root>" + b"<i>x</i>" * 2000 + b"</root>"
    out = render_body(doc, "application/xml", "many.xml")
    assert out.html.count('<div class="json-leaf">') == jsondata.MAX_ITEMS
    assert "… 1,500 more elements" in out.html
    assert "the first 500 children of large elements" in out.html
    deep = b"<a>" * 40 + b"x" + b"</a>" * 40
    assert "16 levels of nesting" in render_body(deep, "application/xml", "deep.xml").html


# --- Per-format ceilings and head previews ------------------------------------


def test_each_format_has_its_own_whole_file_ceiling():
    from agentdrive.rendering.render import (
        DOCUMENT_MAX_BYTES,
        JSON_MAX_BYTES,
        SLIDES_MAX_BYTES,
        WORKBOOK_MAX_BYTES,
        render_ceiling,
    )

    assert render_ceiling("text/markdown", "a.md") == MAX_RENDER_BYTES
    assert render_ceiling("text/x-python", "a.py") == MAX_RENDER_BYTES
    assert render_ceiling("text/csv", "a.csv") == MAX_RENDER_BYTES
    assert render_ceiling("application/json", "a.json") == JSON_MAX_BYTES
    assert render_ceiling("application/x-ndjson", "a.jsonl") == JSON_MAX_BYTES
    assert render_ceiling(XLSX, "a.xlsx") == WORKBOOK_MAX_BYTES
    assert render_ceiling(DOCX, "a.docx") == DOCUMENT_MAX_BYTES
    assert render_ceiling(PPTX, "a.pptx") == SLIDES_MAX_BYTES
    assert JSON_MAX_BYTES > MAX_RENDER_BYTES < WORKBOOK_MAX_BYTES <= DOCUMENT_MAX_BYTES


def test_containers_render_whole_up_to_their_ceiling_and_card_above_it():
    from agentdrive.rendering.render import DOCUMENT_MAX_BYTES, JSON_MAX_BYTES, WORKBOOK_MAX_BYTES

    # Declared sizes above the text cap but under the format's own ceiling
    # render from the whole bytes; the same formats above it are a card even
    # with a source on offer, because a container cannot be read in part.
    docx = _docx(_REPORT)
    assert render_body(docx, DOCX, "r.docx", size_bytes=MAX_RENDER_BYTES + 1).mode == "document"
    assert render_body(b"", DOCX, "r.docx", size_bytes=DOCUMENT_MAX_BYTES + 1,
                       source=BytesSource(docx)).mode == "download"
    book = _workbook(_MODEL)
    assert render_body(book, XLSX, "q.xlsx", size_bytes=MAX_RENDER_BYTES + 1).mode == "sheet"
    assert render_body(b"", XLSX, "q.xlsx", size_bytes=WORKBOOK_MAX_BYTES + 1,
                       source=BytesSource(book)).mode == "download"
    doc = b'{"a": {"b": 1}}'
    over_text_cap = MAX_RENDER_BYTES + 1
    assert render_body(doc, "application/json", "a.json", size_bytes=over_text_cap).mode == "data"
    # A single JSON document over its ceiling is text, so it previews its head
    # as source rather than carding — a tree of a fragment would be a lie.
    over = render_body(b"", "application/json", "a.json", size_bytes=JSON_MAX_BYTES + 1,
                       source=BytesSource(doc))
    assert over.mode == "code" and "download the file for all of it" in over.html


def test_oversized_text_previews_its_head_as_source_and_says_so():
    from agentdrive.rendering.render import HEAD_PREVIEW_BYTES

    lines = [f"[{i:07d}] worker-3 handled request in {i % 97} ms" for i in range(120_000)]
    log = ("\n".join(lines)).encode()
    assert len(log) > MAX_RENDER_BYTES
    source = BytesSource(log)
    out = render_body(b"", "text/plain", "worker.log", size_bytes=len(log), source=source)
    assert out.mode == "code"
    assert source.calls == 1 and source.fetched == HEAD_PREVIEW_BYTES
    assert '<p class="doc-note">Showing the first 1.0 MB of a' in out.html
    assert "[0000000] worker-3" in out.html
    # The line the cut fell in is dropped, never shown half.
    head_text = log[:HEAD_PREVIEW_BYTES].decode()
    last_full = head_text[: head_text.rfind("\n")].splitlines()[-1]
    cut = head_text.splitlines()[-1]
    assert last_full in out.html and (cut == last_full or cut not in out.html)
    assert 'data-view="source"' not in out.html


def test_oversized_markdown_previews_as_source_not_prose():
    md = (b"# Transcript\n\n" + b"- turn: something happened here\n" * 200_000)
    out = render_body(b"", "text/markdown", "t.md", size_bytes=len(md), source=BytesSource(md))
    assert out.mode == "code"  # a fragment is shown as source, not as a cut-off document
    assert "download the file for all of it" in out.html


def test_oversized_json_lines_previews_its_first_records():
    from agentdrive.rendering.render import JSON_MAX_BYTES

    rows = b"".join(b'{"id": %d, "ok": true}\n' % i for i in range(600_000))
    assert len(rows) > JSON_MAX_BYTES
    source = BytesSource(rows)
    out = render_body(b"", "application/x-ndjson", "e.jsonl", size_bytes=len(rows), source=source)
    assert out.mode == "table"
    assert source.calls == 1
    assert "<td>0</td><td>true</td>" in out.html
    assert "Showing the first 1.0 MB of a" in out.html
    assert 'data-view="source"' not in out.html and 'class="doc-modes"' not in out.html


def test_previews_when_large_admits_text_and_parquet_only():
    from agentdrive.rendering.render import previews_when_large

    assert previews_when_large("text/plain", "a.log")
    assert previews_when_large("text/markdown", "a.md")
    assert previews_when_large("application/json", "a.json")
    assert previews_when_large("application/vnd.apache.parquet", "a.parquet")
    assert not previews_when_large(DOCX, "a.docx")
    assert not previews_when_large(PPTX, "a.pptx")
    assert not previews_when_large(XLSX, "a.xlsx")
    assert not previews_when_large("application/zip", "a.zip")
    binary = BytesSource(b"\x00" * 10)
    out = render_body(
        b"", "application/zip", "a.zip", size_bytes=MAX_RENDER_BYTES + 1, source=binary
    )
    assert out.mode == "download" and binary.calls == 0
